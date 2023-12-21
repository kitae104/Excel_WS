import openpyxl
from openpyxl.chart import BubbleChart, Reference

wb = openpyxl.load_workbook("bubble.xlsx")
sheet = wb.active

# 엑셀로 부터 데이터 가져오기
data = [] # 데이터를 저장할 리스트  
for row in sheet.iter_rows(values_only=True):
    data.append([row[0], row[1], row[2]])

chart = BubbleChart()                    # 차트 생성
chart.title = "This is Bubble Chart"     # 차트의 제목 설정

# 차트를 위한 데이터 설정
data_reference = Reference(sheet, min_col=2, min_row=1, max_row=len(data), max_col=3)
categories_reference = Reference(sheet, min_col=1, min_row=2, max_row=len(data), max_col=1)

chart.add_data(data_reference, titles_from_data=True) # titles_from_data=True : 데이터의 제목을 가져옴
chart.set_categories(categories_reference)  # x축의 범주 설정

# 차트를 시트에 추가
sheet.add_chart(chart, "E1")
wb.save("bubble.xlsx")